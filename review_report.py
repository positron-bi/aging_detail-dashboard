"""Read-only comparison of the template and generated report."""
import json
from collections import Counter
from pathlib import Path
import openpyxl

def review(template, output):
    a = openpyxl.load_workbook(template, data_only=True)
    b = openpyxl.load_workbook(output, data_only=True)
    result = {'sheets': [], 'companies': {}}
    for name in a.sheetnames:
        x, y = a[name], b[name]
        result['sheets'].append({'name': name, 'headers_equal': [c.value for c in x[1]] == [c.value for c in y[1]], 'template_rows': x.max_row-1, 'output_rows': y.max_row-1})
    for label, book in [('template', a), ('output', b)]:
        rows = [r for r in book['خلاصه ح.د و وصول'].iter_rows(min_row=2, values_only=True) if r[0] != 'جمع کل']
        result['companies'][label] = dict(Counter(r[0] for r in rows))
        result[label+'_totals'] = {str(i): sum(float(r[i] or 0) for r in rows) for i in range(2,18)}
    sales = list(b['ریز ح.د فروش ماه ۶'].iter_rows(min_row=2,values_only=True))
    receipts = list(b['ریز وصول ح.د'].iter_rows(min_row=2,values_only=True))
    result['nonblank_sale_document_fields'] = sum(bool(r[3]) for r in sales)
    usage = Counter((r[0],r[7]) for r in receipts if r[7])
    result['reused_receipt_counterpart_rows'] = sum(v>1 for v in usage.values())
    result['unallocated_receipts_million_toman'] = sum(max(0,float(r[10] or 0)-float(r[11] or 0)-float(r[12] or 0)) for r in b['خلاصه ح.د و وصول'].iter_rows(min_row=2,values_only=True))
    return result

if __name__ == '__main__':
    import sys
    print(json.dumps(review(sys.argv[1],sys.argv[2]),ensure_ascii=False,indent=2))
